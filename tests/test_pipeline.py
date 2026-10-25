from pathlib import Path
import pandas as pd
from src.generate_data import generate_dataset
from src.validate_data import validate_frames
from src.analyze_experiment import welch_difference, run_analysis
from src.build_executive_artifacts import load_validated_results, build_excel, build_powerpoint
import json
import pytest
from openpyxl import load_workbook
from pptx import Presentation


def write_metadata(data_dir: Path, n_members: int, seed: int = 20261022):
    (data_dir / "generation_metadata.json").write_text(json.dumps({
        "metadata_schema_version": 1, "generator_version": "1.1.0", "seed": seed,
        "n_members": n_members,
        "months": ["2026-01-01", "2026-02-01", "2026-03-01", "2026-04-01", "2026-05-01", "2026-06-01"],
        "synthetic_planted_exposed_member_transaction_intensity_multiplier": 1.15,
        "synthetic_campaign_exposure_probability_among_eligible_treatment_member_months": 0.80,
    }))


def test_generator_schema_and_validation():
    frames = generate_dataset(n_members=1000, seed=1234)
    assert set(frames) == {"raw_members", "raw_member_month", "raw_benefit_events", "raw_market_context"}
    assert len(frames["raw_member_month"]) == 6000
    assert validate_frames(frames) == []
    mm = frames["raw_member_month"].copy()
    mm["month_start"] = pd.to_datetime(mm["month_start"])
    mm = mm.sort_values(["member_id", "month_start"])
    active_start = mm.active_at_start.astype(bool)
    assert not ((~active_start) & (mm.transaction_count.astype(int) > 0)).any()
    assert not (mm.loc[~active_start, "active_at_end"].astype(bool)).any()
    events = frames["raw_benefit_events"]
    exposed = events.campaign_exposed_flag.astype(bool)
    eligible = events.eligible_flag.astype(bool)
    used = events.benefit_used_flag.astype(bool)
    assert not (exposed & ~eligible).any()
    assert not (used & ~eligible).any()
    assert events.groupby("member_id").size().eq(3).all()


def test_reproducibility():
    a = generate_dataset(n_members=500, seed=99)
    b = generate_dataset(n_members=500, seed=99)
    for name in a:
        pd.testing.assert_frame_equal(a[name], b[name])


def test_welch_difference_known_case():
    result = welch_difference([2, 3, 4, 5], [1, 2, 3, 4])
    assert result["difference_treatment_minus_control"] == 1.0
    assert result["n_treatment"] == 4
    assert result["n_control"] == 4


def test_analysis_fails_when_member_event_coverage_is_incomplete(tmp_path):
    frames = generate_dataset(n_members=500, seed=20261022)
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    for name, frame in frames.items():
        frame.to_csv(data_dir / f"{name}.csv", index=False)
    write_metadata(data_dir, 500)
    events = pd.read_csv(data_dir / "raw_benefit_events.csv")
    assigned_ids = set(frames["raw_members"].loc[frames["raw_members"].eligible_for_experiment, "member_id"])
    target = next(iter(assigned_ids))
    events = events[~((events.member_id == target) & (events.month_start == "2026-04-01"))]
    events.to_csv(data_dir / "raw_benefit_events.csv", index=False)
    try:
        run_analysis(data_dir, tmp_path / "outputs")
    except ValueError as exc:
        assert "Input data validation failed" in str(exc)
        assert "benefit-event table expected" in str(exc)
    else:
        raise AssertionError("Analysis must stop before estimating effects when the input data fails validation")


def test_artifact_builder_rejects_old_mock_payload(tmp_path):
    mock = {"primary_estimand": "fake", "analysis_population": "fake", "outcomes": {}, "primary_result": {}}
    path = tmp_path / "mock_results.json"
    path.write_text(json.dumps(mock))
    with pytest.raises(ValueError, match="incomplete/starter/mock"):
        load_validated_results(path)


def test_analysis_results_generate_reopenable_artifacts(tmp_path):
    frames = generate_dataset(n_members=600, seed=20261022)
    data_dir = tmp_path / "data"
    output_dir = tmp_path / "outputs"
    data_dir.mkdir()
    for name, frame in frames.items():
        frame.to_csv(data_dir / f"{name}.csv", index=False)
    write_metadata(data_dir, 600)
    results = run_analysis(data_dir, output_dir)
    assert results["results_schema_version"] == 1
    assert len(results["analysis_run_id"]) == 16
    assert results["primary_result"] == results["outcomes"]["post_contribution"]
    dashboard_results = pd.read_csv(output_dir / "looker_studio_results.csv")
    assert len(dashboard_results) == len(results["outcomes"]) + 1
    primary_row = dashboard_results[(dashboard_results.result_type == "ITT_OUTCOME") & (dashboard_results.outcome == "post_contribution")]
    assert len(primary_row) == 1
    assert bool(primary_row.iloc[0].is_primary)
    assert abs(float(primary_row.iloc[0].effect_estimate) - results["primary_result"]["difference_treatment_minus_control"]) < 1e-12
    assert dashboard_results.analysis_run_id.nunique() == 1
    assert dashboard_results.synthetic_data_only.astype(str).str.lower().eq("true").all()
    artifact_dir = output_dir / "executive_artifacts"
    build_excel(output_dir / "experiment_results.json", artifact_dir / "review.xlsx")
    build_powerpoint(output_dir / "experiment_results.json", artifact_dir / "readout.pptx")
    workbook = load_workbook(artifact_dir / "review.xlsx", read_only=True)
    assert "Input Manifest" in workbook.sheetnames
    assert workbook["Experiment Results"].max_row >= 2
    workbook.close()
    presentation = Presentation(artifact_dir / "readout.pptx")
    assert len(presentation.slides) == 7


def test_bigquery_ddl_is_sandbox_safe_and_financial_sql_has_explicit_aliases():
    repo_root = Path(__file__).resolve().parents[1]
    raw_ddl = (repo_root / "sql" / "02_create_tables.sql").read_text()
    modeled_sql = (repo_root / "sql" / "03_build_analytics.sql").read_text()
    financial_sql = (repo_root / "sql" / "06_financial_scenarios.sql").read_text()
    assert "PARTITION BY month_start" not in raw_ddl
    assert "PARTITION BY month_start" not in modeled_sql
    assert "CLUSTER BY month_start, member_id" in raw_ddl
    assert "CLUSTER BY month_start, member_id" in modeled_sql
    dashboard_sql = (repo_root / "sql" / "07_create_dashboard_views.sql").read_text()
    assert "vw_dashboard_monthly_arm_performance" in dashboard_sql
    assert "vw_dashboard_experiment_arm_summary" in dashboard_sql
    assert "looker_studio_results.csv" in (repo_root / "src" / "analyze_experiment.py").read_text()
    assert (repo_root / "sql" / "08_dashboard_validation.sql").exists()
    assert "DROP TABLE" not in dashboard_sql.upper()
    assert "interval AS (" not in financial_sql
    assert "interval_bounds AS (" in financial_sql
    assert "]) AS scenario_row" in financial_sql
